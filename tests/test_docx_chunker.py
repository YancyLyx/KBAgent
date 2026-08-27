# -*- coding: utf-8 -*-
"""DOCX 分块器单元测试：标题章节、段落、表格序列化"""
import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).parent.parent))

from docx import Document
from app.rag.docx_chunker import DocxChunker


def _make_docx(tmp_path):
    path = tmp_path / "policy.docx"
    doc = Document()
    doc.add_heading("差旅与报销管理办法", level=1)
    doc.add_heading("第一章 总则", level=2)
    doc.add_paragraph("本制度适用于全体员工。")
    doc.add_heading("第二章 差旅标准", level=2)
    doc.add_paragraph("出差住宿按级别执行。")
    table = doc.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "级别"
    table.cell(0, 1).text = "上限"
    table.cell(1, 0).text = "普通员工"
    table.cell(1, 1).text = "400元/晚"
    table.cell(2, 0).text = "管理层"
    table.cell(2, 1).text = "600元/晚"
    doc.save(path)
    return path


def test_docx_heading_table(tmp_path):
    path = _make_docx(tmp_path)
    chunker = DocxChunker("config/rag_config.yaml")
    chunks = chunker.chunk_docx(str(path), metadata={"source": "policy.docx", "tag": "差旅报销"})
    assert chunks, "应产出分块"
    text = "\n".join(c["content"] for c in chunks)
    assert "本制度适用于全体员工" in text and "出差住宿按级别执行" in text, "段落应进入 content"
    assert "普通员工" in text and "400元/晚" in text, "表格内容应被序列化"
    assert any("第一章 总则" in c.get("section", "") for c in chunks), "章节路径应保留"
    assert all(c.get("tag") == "差旅报销" for c in chunks), "metadata 应透传"
    assert all(c.get("chunk_id", "").startswith("policy.docx") for c in chunks)
